import streamlit as st
import pandas as pd
from docx import Document
from datetime import datetime
import io

st.set_page_config(page_title="Generador de Oficios", page_icon="🚗")
st.title("🚗 Generador Automático de Oficios de Comisión")
st.write("Llena los siguientes datos para descargar tu oficio de forma inmediata.")

@st.cache_data
def cargar_datos():
    return pd.read_excel('Información del empleado.xlsx')

try:
    df = cargar_datos()
except Exception as e:
    st.error("❌ No se pudo cargar la base de datos. Verifica los archivos.")
    st.stop()

with st.form("formulario_oficio"):
    num_empleado = st.number_input("1. Número de Empleado:", min_value=1, step=1, format="%d")
    placa_input = st.text_input("2. Placas de la unidad que ocuparás (Ej. HM4036G):").strip().upper()
    lugar_input = st.text_input("3. ¿A qué lugar asistirás y en qué fecha? (Ej. Mercado de Pachuca, 12 de septiembre):").strip()
    motivo_input = st.text_input("4. ¿Cuál es la finalidad de la comisión? (Ej. entregar correspondencia):").strip()
    hora_salida = st.time_input("5. Selecciona tu hora de salida:")
    generar = st.form_submit_button("Generar Oficio")

if generar:
    if not num_empleado or not placa_input or not lugar_input or not motivo_input:
        st.warning("⚠️ Por favor, llena todos los campos antes de generar el oficio.")
    else:
        empleado_data = df[df['No. empleado'] == num_empleado]
        vehiculo_data = df[df['placa'].astype(str).str.upper() == placa_input]
        
        if empleado_data.empty:
            st.error(f"⚠️ No se encontró el número de empleado {num_empleado}.")
        elif vehiculo_data.empty:
            st.error(f"⚠️ No se encontró la placa '{placa_input}'.")
        else:
            datos_emp = empleado_data.iloc[0]
            datos_veh = vehiculo_data.iloc[0]
            
            meses = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
            ahora = datetime.now()
            fecha_larga = f"{ahora.day} de {meses[ahora.month - 1]} de {ahora.year}"
            hora_formateada = hora_salida.strftime("%H:%M")
            
            modelo = str(datos_veh['Modelo'])
            if modelo.endswith('.0'):
                modelo = modelo[:-2]
                
            reemplazos = {
                '[día, mes, año]': fecha_larga, # <-- Acento corregido
                '[Nombre]': str(datos_emp['Nombre']),
                '[Adscripción]': str(datos_emp['Adscripción']),
                '[Puesto]': str(datos_emp['Puesto']),
                '[Nombramiento]': str(datos_emp['Nombramiento']),
                '[RFC]': str(datos_emp['RFC']),
                '[Unidad]': str(datos_veh['Unidad']),
                '[Modelo]': modelo,
                '[Placa]': str(datos_veh['placa']), 
                '[Lugar al que asistirá y Fecha en dia y mes]': lugar_input,
                '[Motivo]': motivo_input,
                '[Horario en tiempo real]': hora_formateada
            }
            
            try:
                doc = Document('OFICIO COMISIÓN 2026_2.docx')
                
                # 1. Buscar en párrafos normales
                for parrafo in doc.paragraphs:
                    for etiqueta, valor_real in reemplazos.items():
                        if etiqueta in parrafo.text:
                            parrafo.text = parrafo.text.replace(etiqueta, valor_real)
                
                # 2. Buscar dentro de las tablas (para las firmas)
                for tabla in doc.tables:
                    for fila in tabla.rows:
                        for celda in fila.cells:
                            for parrafo in celda.paragraphs:
                                for etiqueta, valor_real in reemplazos.items():
                                    if etiqueta in parrafo.text:
                                        parrafo.text = parrafo.text.replace(etiqueta, valor_real)
                
                bio = io.BytesIO()
                doc.save(bio)
                nombre_archivo_salida = f"Oficio_{datos_emp['Nombre'].replace(' ', '_')}_{placa_input}.docx"
                
                st.success(f"✅ ¡Oficio generado exitosamente para {datos_emp['Nombre']}!")
                st.download_button(
                    label="📥 Descargar Documento Word",
                    data=bio.getvalue(),
                    file_name=nombre_archivo_salida,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
            except Exception as e:
                st.error(f"❌ Ocurrió un error al procesar la plantilla: {e}")or(f"❌ Ocurrió un error al procesar la plantilla: {e}")
